"""底稿终表健全性检查：出表前拦截常见结构/刻度问题。

设计目标：即使后续少改代码，每次导出也能自动暴露风险。
规则尽量数据驱动；阈值阈值见本文件顶部常量。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from statistics import median
from typing import Any, Literal

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# 约定（与 normalize_draft / final_tables 对齐；改规则优先改这里）
# ---------------------------------------------------------------------------

Severity = Literal["error", "warning", "info"]

SHEET_ALIASES: dict[str, tuple[str, ...]] = {
    "1整体情况": ("1整体情况", "行业变化"),
    "2管理人总规模": ("2管理人总规模", "总规模带格式", "总规模"),
    "3管理人非货": ("3管理人非货", "非货"),
    "4非货增量": ("4非货增量", "非货增量排名"),
    "5主动权益": ("5主动权益", "主动权益排名"),
    "6权益ETF（含联接）": ("6权益ETF（含联接）", "ETF含联接"),
    "7货币": ("7货币", "货币"),
    "8固收": ("8固收", "固收排名"),
    "9固收+": ("9固收+", "固收+排名"),
    "10 FOF": ("10 FOF", "10FOF", "FOF"),
}

REQUIRED_SHEETS = (
    "1整体情况",
    "2管理人总规模",
    "3管理人非货",
    "4非货增量",
    "5主动权益",
    "6权益ETF（含联接）",
    "7货币",
    "8固收",
    "9固收+",
    "10 FOF",
)

CATEGORY_ORDER = (
    "合计",
    "非货",
    "货币",
    "固收",
    "被动权益",
    "主动权益",
    "固收+",
    "FOF",
)

COMPANY_HEADERS = {"公司", "基金公司", "管理人", "类型"}
TOP_N_HINT = 30
# 增速列：中位数 ≥ 此值视为「百分数」；与 deck_workbook._pct_column_divisor 一致
PCT_POINTS_MEDIAN = 2.0


@dataclass
class Finding:
    severity: Severity
    code: str
    message: str
    sheet: str = ""


@dataclass
class ValidationReport:
    path: str
    findings: list[Finding] = field(default_factory=list)

    @property
    def errors(self) -> list[Finding]:
        return [f for f in self.findings if f.severity == "error"]

    @property
    def warnings(self) -> list[Finding]:
        return [f for f in self.findings if f.severity == "warning"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def add(self, severity: Severity, code: str, message: str, sheet: str = "") -> None:
        self.findings.append(Finding(severity, code, message, sheet=sheet))

    def format_text(self) -> str:
        if not self.findings:
            return f"底稿检查通过：{self.path}（无问题）"
        lines = [f"底稿检查：{self.path}", f"  error={len(self.errors)}  warning={len(self.warnings)}"]
        order = {"error": 0, "warning": 1, "info": 2}
        for f in sorted(self.findings, key=lambda x: (order.get(x.severity, 9), x.sheet, x.code)):
            loc = f"[{f.sheet}] " if f.sheet else ""
            lines.append(f"  {f.severity.upper():7} {f.code}: {loc}{f.message}")
        return "\n".join(lines)

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "ok": self.ok,
            "error_count": len(self.errors),
            "warning_count": len(self.warnings),
            "findings": [
                {"severity": f.severity, "code": f.code, "message": f.message, "sheet": f.sheet}
                for f in self.findings
            ],
        }


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", "").replace(" ", "").strip()


def _resolve_sheet(names: set[str], aliases: tuple[str, ...]) -> str | None:
    for a in aliases:
        if a in names:
            return a
    return None


def _header_row(ws, row: int = 1) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for c in range(1, (ws.max_column or 0) + 1):
        h = _norm(ws.cell(row, c).value)
        if h:
            out.append((c, h))
    return out


def _is_pct_header(h: str) -> bool:
    return "增速" in h or "增幅" in h


def _col_nums(ws, col: int, start_row: int = 2, max_rows: int = 80) -> list[float]:
    vals: list[float] = []
    for r in range(start_row, start_row + max_rows):
        v = ws.cell(r, col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            vals.append(float(v))
    return vals


def _check_sheets(report: ValidationReport, names: set[str]) -> dict[str, str]:
    resolved: dict[str, str] = {}
    for canon in REQUIRED_SHEETS:
        hit = _resolve_sheet(names, SHEET_ALIASES[canon])
        if hit:
            resolved[canon] = hit
        else:
            sev: Severity = "error" if canon in {"1整体情况", "2管理人总规模", "3管理人非货"} else "warning"
            report.add(
                sev,
                "missing_sheet",
                f"缺少 sheet「{canon}」（别名：{' / '.join(SHEET_ALIASES[canon])}）",
            )
    return resolved


def _check_industry(report: ValidationReport, ws) -> None:
    sheet = ws.title
    headers = _header_row(ws)
    dates = [h for _, h in headers if re.match(r"^\d{8}", h)]
    if len(dates) < 2:
        report.add(
            "error",
            "industry_dates",
            f"第1行需含至少 2 个 YYYYMMDD 日期列（期末[+上季末]+年初；Q1 可仅期末+年初），当前识别到 {dates or '无'}",
            sheet,
        )
    elif len(dates) == 2:
        report.add(
            "info",
            "industry_dates_q1",
            f"日期列（Q1 双列口径）：{' → '.join(dates)}（期末→年初；上季末按年初处理）",
            sheet,
        )
    elif len(dates) > 3:
        report.add(
            "warning",
            "industry_dates_extra",
            f"识别到超过 3 个日期列：{dates}，请确认前三个为期末/上季末/年初",
            sheet,
        )
    else:
        report.add("info", "industry_dates_ok", f"日期列：{' → '.join(dates[:3])}", sheet)

    # 类型列
    first = _norm(ws.cell(1, 1).value)
    if first and first != "类型" and not re.match(r"^\d{8}", first):
        report.add("warning", "industry_type_header", f"A1 建议为「类型」，当前为「{first}」", sheet)
    elif not first:
        report.add("info", "industry_type_blank", "A1 为空，导出时会自动补「类型」", sheet)

    # 品类
    labels: list[str] = []
    blank_after: list[str] = []
    seen_blank = False
    for r in range(2, min((ws.max_row or 1) + 1, 40)):
        lab = _norm(ws.cell(r, 1).value)
        if not lab:
            # 空行
            if labels and not seen_blank:
                seen_blank = True
            continue
        if lab in CATEGORY_ORDER:
            if seen_blank and lab in CATEGORY_ORDER:
                blank_after.append(lab)
            labels.append(lab)

    missing = [c for c in CATEGORY_ORDER if c not in labels]
    if missing:
        report.add(
            "error" if {"合计", "非货", "货币"} & set(missing) else "warning",
            "industry_categories",
            f"主品类缺失：{', '.join(missing)}（需要：{' → '.join(CATEGORY_ORDER)}）",
            sheet,
        )
    if blank_after:
        report.add(
            "warning",
            "industry_blank_row",
            f"主品类之间出现空行，可能导致截断：空行后仍有 {', '.join(blank_after)}",
            sheet,
        )

    # 增速刻度
    for c, h in headers:
        if not _is_pct_header(h):
            continue
        _check_pct_column(report, ws, c, h, sheet)


def _check_pct_column(report: ValidationReport, ws, col: int, header: str, sheet: str) -> None:
    nums = _col_nums(ws, col)
    if len(nums) < 3:
        return
    abs_vals = [abs(x) for x in nums]
    med = median(abs_vals)
    # (0,1) 非整数：既可能是 Excel 小数(0.02=2%)，也可能是百分数里很小的增速(0.84=0.84%)
    # 不能单凭 (0,1) 就断定「混用」。以列中位数为准：
    # - median≥2 → 列是百分数；0.84 只是小增速，不报错
    # - median<1 且仍有一批 |v|≥2 → 才像真混用
    clear_frac = sum(1 for v in abs_vals if 0 < v < 1)
    clear_points = sum(1 for v in abs_vals if v >= PCT_POINTS_MEDIAN)
    n = len(abs_vals)
    need = max(2, int(0.15 * n))

    if med >= PCT_POINTS_MEDIAN:
        report.add(
            "info",
            "pct_scale",
            f"列「{header}」判定为 百分数(2=2%)（median={med:.3f}；"
            f"含 {clear_frac} 个 |v|<1 的小增速，按百分数保留）",
            sheet,
        )
        return

    if med < 1.0 and clear_points >= need:
        report.add(
            "warning",
            "pct_mixed_scale",
            (
                f"列「{header}」整体像 Excel 小数（median={med:.3f}），"
                f"但又有 {clear_points} 个 |v|≥{PCT_POINTS_MEDIAN:g} 的百分数，"
                f"可能导致 1%/2% 显示成 100%/200%。建议统一为百分数或统一为 Excel 小数。"
            ),
            sheet,
        )
        return

    # 中位数落在模糊带：两边证据都强 → 提示，默认不拦生成（底稿通常可信）
    if clear_frac >= need and clear_points >= need:
        report.add(
            "warning",
            "pct_mixed_scale",
            (
                f"列「{header}」疑似同时含小数（如 0.02）与百分数（如 5），"
                f"可能导致 1%/2% 显示成 100%/200%。建议统一为百分数或统一为 Excel 小数。"
                f"（median={med:.3f}, 明确小数={clear_frac}, 明确百分数={clear_points}）"
            ),
            sheet,
        )
        return

    style = "百分数(2=2%)" if med >= PCT_POINTS_MEDIAN else "Excel小数(0.02=2%)"
    report.add("info", "pct_scale", f"列「{header}」判定为 {style}（median={med:.3f}）", sheet)


def _check_rank_sheet(report: ValidationReport, ws, *, kind: str = "rank") -> None:
    sheet = ws.title
    headers = _header_row(ws)
    if not headers:
        report.add("error", "empty_headers", "第1行无表头", sheet)
        return

    # 表头空洞：前 N 列中间出现空表头会截断
    max_c = max(c for c, _ in headers)
    holes = [c for c in range(1, max_c + 1) if not _norm(ws.cell(1, c).value)]
    if holes and kind != "etf":
        # 仅当空洞落在已用区间内
        report.add(
            "warning",
            "header_gap",
            f"表头第 {holes[:8]} 列为空，连续读表会在此截断",
            sheet,
        )

    has_company = any(h in COMPANY_HEADERS or h == "公司" for _, h in headers)
    if kind != "industry" and not has_company and kind != "etf":
        # ETF / 行业另论；排名表必须有公司列
        if kind == "rank":
            report.add("error", "missing_company_col", "缺少公司列（公司/基金公司/管理人）", sheet)

    # 排名列数值行数
    data_rows = 0
    for r in range(2, min((ws.max_row or 1) + 1, 80)):
        v = ws.cell(r, 1).value
        try:
            float(v)
            data_rows += 1
        except (TypeError, ValueError):
            if data_rows >= 5:
                break
    if kind == "rank":
        if data_rows < 10:
            report.add("warning", "few_rows", f"可识别排名行仅 {data_rows} 行（建议 ≥{TOP_N_HINT}）", sheet)
        elif data_rows < TOP_N_HINT:
            report.add("info", "rows_lt_topn", f"排名行 {data_rows} < Top{TOP_N_HINT}，导出将少于 30 行", sheet)

    for c, h in headers:
        if _is_pct_header(h):
            _check_pct_column(report, ws, c, h, sheet)

    # 规模相关列提示（仅非货主表）
    if kind == "rank" and ("管理人非货" in sheet or sheet in {"3管理人非货", "非货"}):
        aum_bits = [h for _, h in headers if h in {
            "主动权益", "被动权益", "固收+", "固收", "FOF",
            "主动权益规模", "被动权益规模", "固收+规模", "固收规模", "FOF规模",
        }]
        if len(aum_bits) < 3:
            report.add(
                "warning",
                "non_money_breakdown",
                "非货分类规模列偏少（建议含 主动权益/被动权益/固收+/固收/FOF，可带或不带「规模」后缀）",
                sheet,
            )


def _check_etf(report: ValidationReport, ws) -> None:
    sheet = ws.title
    headers = _header_row(ws)
    if not headers:
        report.add("error", "empty_headers", "第1行无表头", sheet)
        return

    right = None
    for c, h in headers:
        if "权益ETF" in h and "排名" in h:
            right = (c, h)
            break
    if right is None:
        report.add(
            "warning",
            "etf_right_missing",
            "未找到右表起点（表头需同时含「权益ETF」与「排名」）；将跳过 06_权益ETF含联接",
            sheet,
        )
    else:
        report.add("info", "etf_right_ok", f"右表起点列 {right[0]}：「{right[1]}」", sheet)

    # 左表：第 1 列起连续表头，至空列或右表起点（宽度随底稿变化，如 Q3=13）
    left_cols = 0
    max_scan = (right[0] - 1) if right else max((c for c, _ in headers), default=0)
    for c in range(1, max_scan + 1):
        h = _norm(ws.cell(1, c).value)
        if not h or ("权益ETF" in h and "排名" in h):
            break
        left_cols = c
    if left_cols:
        report.add(
            "info",
            "etf_left_cols",
            f"左表连续 {left_cols} 列（动态宽度）",
            sheet,
        )
    left_headers = [h for c, h in headers if c <= left_cols] if left_cols else []
    if not any("规模" in h or h in COMPANY_HEADERS for h in left_headers):
        report.add("warning", "etf_left_sparse", "左表缺少规模/公司类表头", sheet)

    for c, h in headers:
        if _is_pct_header(h):
            _check_pct_column(report, ws, c, h, sheet)


def validate_draft_xlsx(xlsx_path: str | Path) -> ValidationReport:
    """检查底稿终表结构与增速刻度；不修改文件。"""
    path = Path(xlsx_path)
    report = ValidationReport(path=str(path))
    if not path.exists():
        report.add("error", "file_missing", f"文件不存在：{path}")
        return report

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        report.add("error", "open_failed", f"无法打开工作簿：{exc}")
        return report

    try:
        names = set(wb.sheetnames)
        resolved = _check_sheets(report, names)

        if "1整体情况" in resolved:
            _check_industry(report, wb[resolved["1整体情况"]])

        for canon in (
            "2管理人总规模",
            "3管理人非货",
            "4非货增量",
            "5主动权益",
            "7货币",
            "8固收",
            "9固收+",
            "10 FOF",
        ):
            if canon in resolved:
                _check_rank_sheet(report, wb[resolved[canon]], kind="rank")

        if "6权益ETF（含联接）" in resolved:
            _check_etf(report, wb[resolved["6权益ETF（含联接）"]])
    finally:
        wb.close()

    if report.ok and not report.warnings:
        report.add("info", "all_clear", "结构与增速刻度检查通过")
    return report
