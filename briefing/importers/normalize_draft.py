"""将任意期底稿整理为「案例终表」口径，供后续 load_final_tables / 出表使用。

对齐 25Q4 简报底稿约定：
- 行业表：首列表头「类型」；仅合计~FOF 主品类，固定顺序
- 排名表：保留 TopN（默认 30）；关注公司若在 TopN 外仍追加一行
- ETF：左表按连续表头动态列宽（可 >10，如 Q3 含单季列）；右表定位后写到 ≥20 列；左右均 TopN（含关注公司兜底）
- 公司列统一简称（去掉「基金」后缀）；常见表头别名归一
"""

from __future__ import annotations

import re
import warnings
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from briefing.excel.table_layout import format_header_wrap

# 案例终表 sheet 名（输出侧）
_SHEET_OUT = {
    "1整体情况": ("1整体情况", "行业变化"),
    "2管理人总规模": ("2管理人总规模", "总规模带格式", "总规模"),
    "3管理人非货": ("3管理人非货", "非货"),
    "4非货增量": ("4非货增量", "非货增量排名"),
    "5主动权益": ("5主动权益", "主动权益排名"),
    "6权益ETF（含联接）": ("6权益ETF（含联接）", "ETF含联接"),
    "7货币": ("7货币", "货币"),
    "8固收": ("8固收", "固收排名"),
    "9固收+": ("9固收+", "固收+排名"),
    "10 FOF": ("10 FOF", "10FOF", "FOF"),
}

_CATEGORY_ORDER = (
    "合计",
    "非货",
    "货币",
    "固收",
    "被动权益",
    "主动权益",
    "固收+",
    "FOF",
)

_TOP_N_DEFAULT = 30

# 公司列可能的表头
_COMPANY_HEADERS = {"公司", "基金公司", "管理人"}


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", "").replace(" ", "").strip()


def _resolve_in_sheet(wb, aliases: tuple[str, ...]) -> str | None:
    names = set(wb.sheetnames)
    for a in aliases:
        if a in names:
            return a
    return None


def _is_blank_row(ws, row: int, max_col: int) -> bool:
    for c in range(1, max_col + 1):
        v = ws.cell(row, c).value
        if v is not None and str(v).strip() != "":
            return False
    return True


def _short_company(v: Any) -> Any:
    if v is None:
        return v
    s = str(v).strip()
    if s.endswith("基金") and len(s) > 2:
        return s[:-2]
    return s


def _find_company_col(headers: list[str]) -> int | None:
    for i, h in enumerate(headers):
        if _norm(h) in _COMPANY_HEADERS:
            return i
    return None


def _header_aliases(h: Any, *, ytd_tag: str, q_tag: str) -> Any:
    """表头别名 → 案例风格（尽量保留原日期数字列）。"""
    if h is None:
        return h
    raw = str(h)
    n = _norm(h)
    mapping = {
        "基金公司": "公司",
        "管理人": "公司",
        "非货ETF（不剔联接）排名": "非货ETF排名",
        "非货ETF(不剔联接)排名": "非货ETF排名",
        "权益ETF（国内外不踢联接）排名": "权益ETF\n排名",
        "权益ETF(国内外不踢联接)排名": "权益ETF\n排名",
        "YTD增量": f"{ytd_tag}增量",
        "YTD规模增量": f"{ytd_tag}规模增量",
        "YTD规模增幅%": f"{ytd_tag}规模增幅%",
        "YTD排名变化": f"{ytd_tag}排名变化",
        "YTD增速%": f"{ytd_tag}增速",
        "YTD增速": f"{ytd_tag}增速",
        "YTD新发": f"{ytd_tag}新发",
        "YTD净值变化": f"{ytd_tag}净值变化",
        "YTD持营": f"{ytd_tag}持营",
        "季度增量": f"{q_tag}增量" if q_tag else "季度增量",
        "季度规模增量": f"{q_tag}规模增量" if q_tag else "季度规模增量",
        "季度规模增幅%": f"{q_tag}规模增幅%" if q_tag else "季度规模增幅%",
        "季度排名变化": f"{q_tag}排名变化" if q_tag else "季度排名变化",
        "季度增速%": f"{q_tag}增速" if q_tag else "季度增速",
        "季度增速": f"{q_tag}增速" if q_tag else "季度增速",
        "季度新发": f"{q_tag}新发" if q_tag else "季度新发",
        "季度净值影响": f"{q_tag}净值影响" if q_tag else "季度净值影响",
        "季度持营": f"{q_tag}持营" if q_tag else "季度持营",
    }
    if n in mapping:
        return mapping[n]
    # 带换行的排名类
    return format_header_wrap(raw)


def _detect_ytd_q_tags(ws_industry) -> tuple[str, str]:
    """从行业表日期列推断 ytd/quarter 口语前缀（与 PeriodProfile 对齐）。"""
    from briefing.period_profile import resolve_period_profile

    dates: list[str] = []
    for c in range(1, min(8, (ws_industry.max_column or 1) + 1)):
        h = _norm(ws_industry.cell(1, c).value)
        m = re.match(r"^(\d{8})", h)
        if m:
            dates.append(m.group(1))
    if len(dates) < 1:
        return "YTD", "季度"
    cur = dates[0]
    profile = resolve_period_profile("", "", cur)
    return profile.ytd_tag, profile.quarter_tag


def _company_matches(name: Any, focus_company: str, focus_short: str) -> bool:
    if name is None:
        return False
    n = str(name).replace("基金", "").strip()
    short = (focus_short or focus_company or "").replace("基金", "").strip()
    full = (focus_company or "").strip()
    if short and short in n:
        return True
    if full and full in str(name):
        return True
    return False


def _copy_top_n_block(
    src_ws,
    dst_ws,
    *,
    src_start_col: int,
    n_cols: int,
    dst_start_col: int,
    top_n: int,
    ytd_tag: str,
    q_tag: str,
    strip_company: bool = True,
    focus_company: str = "示例基金",
    focus_short: str = "示例",
) -> int:
    """复制表头 + TopN；若关注公司排名在 TopN 外则追加一行（如固收第32的示例）。"""
    headers = []
    for i in range(n_cols):
        h = src_ws.cell(1, src_start_col + i).value
        headers.append(h)
        dst_ws.cell(1, dst_start_col + i, _header_aliases(h, ytd_tag=ytd_tag, q_tag=q_tag))

    company_i = _find_company_col([_norm(h) for h in headers])

    def _write_src_row(src_r: int, out_r: int) -> None:
        for i in range(n_cols):
            v = src_ws.cell(src_r, src_start_col + i).value
            if strip_company and company_i is not None and i == company_i:
                v = _short_company(v)
            dst_ws.cell(out_r, dst_start_col + i, v)

    candidates: list[tuple[float, int]] = []  # (rank, src_row)
    focus_src_row: int | None = None
    max_scan = min(src_ws.max_row or 1, 500)
    blank_streak = 0
    for r in range(2, max_scan + 1):
        first = src_ws.cell(r, src_start_col).value
        if first is None or str(first).strip() == "":
            blank_streak += 1
            if candidates and blank_streak >= 2:
                break
            continue
        blank_streak = 0
        try:
            rank = float(first)
        except (TypeError, ValueError):
            continue
        candidates.append((rank, r))
        co_val = (
            src_ws.cell(r, src_start_col + company_i).value
            if company_i is not None
            else None
        )
        if _company_matches(co_val, focus_company, focus_short):
            focus_src_row = r

    written = 0
    focus_included = False
    for rank, src_r in candidates:
        if rank > top_n:
            continue
        if written >= top_n:
            break
        written += 1
        _write_src_row(src_r, written + 1)
        if focus_src_row == src_r:
            focus_included = True

    if focus_src_row is not None and not focus_included:
        written += 1
        _write_src_row(focus_src_row, written + 1)

    return written


def _find_etf_right_start(ws) -> int | None:
    for c in range(1, (ws.max_column or 1) + 1):
        h = _norm(ws.cell(1, c).value)
        if "权益ETF" in h and "排名" in h:
            return c
    return None


def _normalize_industry(src_ws, dst_ws, *, ytd_tag: str, q_tag: str) -> None:
    # 读表头：允许首格为空
    headers: list[Any] = []
    start_c = 1
    if not _norm(src_ws.cell(1, 1).value):
        headers.append("类型")
        start_c = 2
    for c in range(start_c, (src_ws.max_column or 1) + 1):
        h = src_ws.cell(1, c).value
        if h is None or _norm(h) == "":
            break
        headers.append(_header_aliases(h, ytd_tag=ytd_tag, q_tag=q_tag))

    for i, h in enumerate(headers, 1):
        dst_ws.cell(1, i, h if i > 1 or h == "类型" else "类型")
    dst_ws.cell(1, 1, "类型")

    # 收集品类行
    by_label: dict[str, list[Any]] = {}
    n_cols = len(headers)
    for r in range(2, (src_ws.max_row or 1) + 1):
        label = src_ws.cell(r, 1).value
        if label is None or str(label).strip() == "":
            if by_label:
                break
            continue
        label_s = str(label).strip()
        if label_s not in _CATEGORY_ORDER:
            continue
        vals = [src_ws.cell(r, c).value for c in range(1, n_cols + 1)]
        by_label[label_s] = vals

    out_r = 2
    for name in _CATEGORY_ORDER:
        if name not in by_label:
            continue
        vals = by_label[name]
        for c, v in enumerate(vals, 1):
            dst_ws.cell(out_r, c, name if c == 1 else v)
        out_r += 1


def _normalize_rank_sheet(
    src_ws,
    dst_ws,
    *,
    top_n: int,
    ytd_tag: str,
    q_tag: str,
    focus_company: str = "示例基金",
    focus_short: str = "示例",
) -> None:
    n_cols = 0
    for c in range(1, (src_ws.max_column or 1) + 1):
        if src_ws.cell(1, c).value is None or _norm(src_ws.cell(1, c).value) == "":
            break
        n_cols = c
    if n_cols <= 0:
        return
    _copy_top_n_block(
        src_ws,
        dst_ws,
        src_start_col=1,
        n_cols=n_cols,
        dst_start_col=1,
        top_n=top_n,
        ytd_tag=ytd_tag,
        q_tag=q_tag,
        focus_company=focus_company,
        focus_short=focus_short,
    )


def _normalize_etf_sheet(
    src_ws,
    dst_ws,
    *,
    top_n: int,
    ytd_tag: str,
    q_tag: str,
    focus_company: str = "示例基金",
    focus_short: str = "示例",
) -> None:
    """左表：连续表头动态列宽；右表写到 max(20, 左表列数+2)，避免与左表重叠。"""
    right_start = _find_etf_right_start(src_ws)
    left_cols = 0
    scan_to = (right_start - 1) if right_start else (src_ws.max_column or 1)
    for c in range(1, scan_to + 1):
        h = _norm(src_ws.cell(1, c).value)
        if not h or ("权益ETF" in h and "排名" in h):
            break
        left_cols = c
    if left_cols <= 0:
        left_cols = 10  # 极端缺表头时仍占位，避免空左表

    _copy_top_n_block(
        src_ws,
        dst_ws,
        src_start_col=1,
        n_cols=left_cols,
        dst_start_col=1,
        top_n=top_n,
        ytd_tag=ytd_tag,
        q_tag=q_tag,
        focus_company=focus_company,
        focus_short=focus_short,
    )

    if right_start is None:
        return
    right_cols = 0
    for c in range(right_start, (src_ws.max_column or 1) + 1):
        h = _norm(src_ws.cell(1, c).value)
        if not h:
            break
        right_cols += 1
    if right_cols <= 0:
        return
    # 案例布局右表常从 20 列起；左表更宽时（如 Q3=13）向右顺延，避免覆盖
    right_dst = max(20, left_cols + 2)
    _copy_top_n_block(
        src_ws,
        dst_ws,
        src_start_col=right_start,
        n_cols=right_cols,
        dst_start_col=right_dst,
        top_n=top_n,
        ytd_tag=ytd_tag,
        q_tag=q_tag,
        focus_company=focus_company,
        focus_short=focus_short,
    )


def normalize_draft_xlsx(
    src_xlsx: str | Path,
    dst_xlsx: str | Path,
    *,
    top_n: int = _TOP_N_DEFAULT,
    focus_company: str = "示例基金",
    focus_short: str = "示例",
) -> Path:
    """
    读入原始底稿，写出案例口径终表 xlsx。
    TopN 之外若含关注公司，仍追加一行，避免 PPT 读表报 focus_missing。
    """
    src = Path(src_xlsx)
    dst = Path(dst_xlsx)
    if not src.exists():
        raise FileNotFoundError(src)
    dst.parent.mkdir(parents=True, exist_ok=True)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        src_wb = load_workbook(src, data_only=True)

    ind_name = _resolve_in_sheet(src_wb, _SHEET_OUT["1整体情况"])
    ytd_tag, q_tag = "YTD", "季度"
    if ind_name:
        ytd_tag, q_tag = _detect_ytd_q_tags(src_wb[ind_name])

    out = Workbook()
    out.remove(out.active)

    for out_name, aliases in _SHEET_OUT.items():
        in_name = _resolve_in_sheet(src_wb, aliases)
        if not in_name:
            continue
        src_ws = src_wb[in_name]
        dst_ws = out.create_sheet(out_name)

        if out_name == "1整体情况":
            _normalize_industry(src_ws, dst_ws, ytd_tag=ytd_tag, q_tag=q_tag)
        elif out_name == "6权益ETF（含联接）":
            _normalize_etf_sheet(
                src_ws,
                dst_ws,
                top_n=top_n,
                ytd_tag=ytd_tag,
                q_tag=q_tag,
                focus_company=focus_company,
                focus_short=focus_short,
            )
        else:
            _normalize_rank_sheet(
                src_ws,
                dst_ws,
                top_n=top_n,
                ytd_tag=ytd_tag,
                q_tag=q_tag,
                focus_company=focus_company,
                focus_short=focus_short,
            )

    if not out.sheetnames:
        raise ValueError(f"未能从底稿识别任何终表 sheet: {src}")

    out.save(dst)
    return dst
