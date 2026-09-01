"""全册 PPT：默认只灌各页文字；表图可由人工粘贴。

有底稿/终表时，文字从终表读数生成（与贴图表口径一致）。
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

from briefing.data_loader import load_config, load_fund_data
from briefing.deck.data_tables import (
    build_business_rows,
    build_category_rows,
    build_increment_rows,
    build_non_money_rows,
    build_total_rows,
    cols_business,
    cols_category,
    cols_increment,
    cols_non_money,
    cols_total,
)
from briefing.deck.narratives import build_all_narratives
from briefing.deck.sheet_narratives import build_narratives_from_tables
from briefing.excel.deck_workbook import build_deck_workbook
from briefing.excel.sheet_export import excel_available, export_sheets_via_excel
from briefing.excel.styled_table import ColumnSpec, render_styled_table
from briefing.render.pptx_filler import (
    _find_shape,
    load_slide_map,
    remove_picture,
    replace_narrative_paragraphs,
    replace_picture_fit,
)

# 终表 sheet 名 → PPT 贴图 key（仅 paste_tables=True 时使用）
_SHEET_IMAGE_KEY = {
    "01_行业变化": "category_summary",
    "02_总规模": "total_ranking_top30",
    "03_非货": "non_money_ranking",
    "04_非货增量": "increment_overview",
    "05_主动权益": "active_equity",
    "06_非货ETF含联接": "etf_with_link",
    "06_权益ETF含联接": "etf_no_link",
    "06_ETF含联接": "etf_with_link",
    "06_ETF不含联接": "etf_no_link",
    "07_货币": "money",
    "08_固收": "fixed_income",
    "09_固收+": "fixed_income_plus",
    "10_FOF": "fof",
}


def _fill_slide_text(
    slide,
    shape_name: str,
    para_map: dict[int, str],
    font_size: float = 12,
    *,
    bold_token: str | None = None,
) -> None:
    shape = _find_shape(slide, shape_name)
    replace_narrative_paragraphs(
        shape, para_map, font_size_pt=font_size, bold_token=bold_token
    )


def _business_dict_rows(rankings) -> list[dict]:
    return [
        {
            "rank": r.rank,
            "company": r.company,
            "aum": r.aum,
            "increment": r.increment,
            "growth_pct": r.growth_pct,
            "new_issue": r.new_issue,
            "nav_change": r.nav_change,
            "holding_sales": r.holding_sales,
            "rank_change": r.rank_change,
            "q_increment": getattr(r, "q_increment", 0),
            "q_growth_pct": getattr(r, "q_growth_pct", 0),
            "rank_change_q": getattr(r, "rank_change_q", 0),
        }
        for r in rankings
    ]


def _default_xlsx_path(output_pptx: Path) -> Path:
    stem = output_pptx.stem
    for suffix in ("_briefing", "_full_deck"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    return output_pptx.with_name(f"{stem}_tables.xlsx")


def _is_tables_workbook(path: Path) -> bool:
    """判断是否为导出的终表工作簿（01_行业变化…）。"""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=False)
        names = set(wb.sheetnames)
        wb.close()
    except Exception:
        return False
    return "01_行业变化" in names or "02_总规模" in names


def _render_table_images(
    excel_sheets: list[tuple[str, list[ColumnSpec], list[dict]]],
    work_dir: Path,
    focus: str,
) -> dict[str, Path]:
    images: dict[str, Path] = {}
    for title, cols, rows in excel_sheets:
        key = _SHEET_IMAGE_KEY.get(title)
        if not key:
            continue
        focus_for = "__none__" if title.startswith("01_") else focus
        safe = title.replace("+", "plus").replace("/", "_")
        images[key] = render_styled_table(
            cols,
            rows,
            work_dir / f"{safe}.png",
            focus_company=focus_for,
            autofit=True,
        )
    return images


def _images_from_excel_export(
    xlsx_path: Path,
    excel_sheets: list[tuple[str, list[ColumnSpec], list[dict]]],
    work_dir: Path,
) -> dict[str, Path]:
    if not excel_available():
        return {}
    names = [t[:31] for t, _, _ in excel_sheets]
    exported = export_sheets_via_excel(xlsx_path, names, work_dir / "_excel_shots")
    if not exported:
        return {}
    images: dict[str, Path] = {}
    for title, _, _ in excel_sheets:
        key = _SHEET_IMAGE_KEY.get(title)
        if not key:
            continue
        png = exported.get(title[:31])
        if png and png.exists():
            images[key] = png
    return images


def _build_sheets_from_csv(df, config, narratives: dict):
    c_cat, c_tot, c_nm, c_inc, c_biz = (
        cols_category(config),
        cols_total(config),
        cols_non_money(config),
        cols_increment(config),
        cols_business(config),
    )
    excel_sheets: list[tuple[str, list[ColumnSpec], list[dict]]] = []
    excel_sheets.append(("01_行业变化", c_cat, build_category_rows(df, config)))
    excel_sheets.append(("02_总规模", c_tot, build_total_rows(df, config)))
    excel_sheets.append(("03_非货", c_nm, build_non_money_rows(df, config)))
    excel_sheets.append(("04_非货增量", c_inc, build_increment_rows(df, config)))
    excel_sheets.append(
        ("05_主动权益", c_biz, build_business_rows(df, config, "active_equity"))
    )
    excel_sheets.append(
        (
            "06_ETF含联接",
            c_biz,
            build_business_rows(df, config, "passive_equity_etf_with_link"),
        )
    )
    excel_sheets.append(
        (
            "06_ETF不含联接",
            c_biz,
            build_business_rows(df, config, "passive_equity_etf_no_link"),
        )
    )
    for sheet_title, cat in [
        ("07_货币", "money"),
        ("08_固收", "fixed_income"),
        ("09_固收+", "fixed_income_plus"),
        ("10_FOF", "fof"),
    ]:
        rows = build_business_rows(df, config, cat)
        if cat == "fixed_income" and narratives.get("_fi_table_rankings"):
            rows = _business_dict_rows(narratives["_fi_table_rankings"])
        if cat == "fixed_income_plus" and narratives.get("_fip_table_rankings"):
            rows = _business_dict_rows(narratives["_fip_table_rankings"])
        excel_sheets.append((sheet_title, c_biz, rows))
    return excel_sheets


def run_full_deck(
    config_path: str | Path = "config/report_26h1.yaml",
    data_dir: str | Path = "data/sample_25h1",
    template_path: str | Path = "templates/25Q4_template.pptx",
    slide_map_path: str | Path = "config/slide_map_q4.yaml",
    output_pptx: str | Path = "output/26H1_briefing.pptx",
    output_xlsx: str | Path | None = None,
    work_dir: str | Path = "output/_full_deck",
    draft_xlsx: str | Path | None = None,
    tables_xlsx: str | Path | None = None,
    *,
    paste_tables: bool = False,
    strict: bool = False,
) -> dict:
    """
    生成全册 PPT。

    - 默认只写入文字（paste_tables=False），表图留给人工粘贴。
    - 有 draft_xlsx / tables_xlsx 时，文字从终表读数，避免与 CSV 双算不一致。
    - strict=True 且读表存在 error 时抛出 ValueError。
    """
    config = load_config(config_path)
    df = load_fund_data(data_dir)
    work_dir = Path(work_dir)
    if work_dir.exists():
        shutil.rmtree(work_dir, ignore_errors=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    slide_map = load_slide_map(slide_map_path)
    focus = config.focus_company_short or config.focus_company

    excel_sheets: list[tuple[str, list[ColumnSpec], list[dict]]] = []
    draft_path = Path(draft_xlsx) if draft_xlsx else None
    tables_path = Path(tables_xlsx) if tables_xlsx else None
    narrative_report = ""
    narrative_source = "csv"
    csv_narratives: dict | None = None

    if draft_path and draft_path.exists():
        from briefing.deck.final_tables import load_final_tables
        from briefing.importers.normalize_draft import normalize_draft_xlsx

        normalized = work_dir / "draft_normalized.xlsx"
        normalize_draft_xlsx(
            draft_path,
            normalized,
            focus_company=config.focus_company,
            focus_short=config.focus_company_short or focus,
        )
        excel_sheets = load_final_tables(
            normalized,
            focus_company=config.focus_company,
            focus_short=config.focus_company_short or focus,
        )
    elif tables_path and tables_path.exists():
        # 直接使用已导出的终表；不再从 CSV 成文
        excel_sheets = []
    else:
        csv_narratives = build_all_narratives(df, config)
        excel_sheets = _build_sheets_from_csv(df, config, csv_narratives)

    output_pptx = Path(output_pptx)
    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    if output_xlsx:
        xlsx_path = Path(output_xlsx)
    elif tables_path and tables_path.exists():
        xlsx_path = tables_path
    else:
        xlsx_path = _default_xlsx_path(output_pptx)

    if excel_sheets:
        xlsx_path = build_deck_workbook(
            excel_sheets,
            xlsx_path,
            period_label=config.period_label,
            focus_company=focus,
        )
    elif tables_path and tables_path.exists() and xlsx_path.resolve() != tables_path.resolve():
        xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(tables_path, xlsx_path)

    # 关键：文字必须从「最终写出的 tables.xlsx」读入，与人工贴表同一文件
    if xlsx_path.exists() and _is_tables_workbook(xlsx_path):
        narratives, book = build_narratives_from_tables(config, tables_xlsx=str(xlsx_path))
        narrative_report = book.format_issues()
        narrative_source = "tables"
        if strict and book.errors:
            raise ValueError(f"终表读数失败（strict）：\n{narrative_report}")
    elif draft_path and draft_path.exists():
        raise ValueError(
            f"已提供底稿但未能写出可读终表：{xlsx_path}。请先 export-tables 或检查底稿。"
        )
    else:
        narratives = csv_narratives or build_all_narratives(df, config)
        narrative_report = (
            "警告：未提供底稿/终表，文字来自 CSV 指标（可能与 Excel 终表不一致）。"
            "请使用 --draft-xlsx 或 --tables-xlsx。"
        )
        narrative_source = "csv"

    images: dict[str, Path] = {}
    used_excel = False
    if paste_tables and excel_sheets:
        images = _images_from_excel_export(xlsx_path, excel_sheets, work_dir)
        used_excel = bool(images)
        if len(images) < len([t for t, _, _ in excel_sheets if t in _SHEET_IMAGE_KEY]):
            pil = _render_table_images(excel_sheets, work_dir, focus)
            for k, v in pil.items():
                images.setdefault(k, v)

    shutil.copy(template_path, output_pptx)
    prs = Presentation(str(output_pptx))

    for slide_cfg in slide_map.get("slides", []):
        idx = slide_cfg["index"]
        sid = slide_cfg["id"]
        slide = prs.slides[idx]
        payload = narratives.get(sid, {})

        if "narrative_shape" in slide_cfg and "paragraphs" in slide_cfg:
            para_map = {}
            for para_idx, key in slide_cfg["paragraphs"].items():
                if key in payload:
                    para_map[int(para_idx)] = payload[key]
            if para_map:
                _fill_slide_text(
                    slide,
                    slide_cfg["narrative_shape"],
                    para_map,
                    font_size=11,
                    bold_token=None,
                )

        for extra in slide_cfg.get("extra_text", []):
            if extra["key"] in payload:
                _fill_slide_text(
                    slide,
                    extra["shape"],
                    {0: payload[extra["key"]]},
                    font_size=11,
                    bold_token=None,
                )

        if "footnote_shape" in slide_cfg:
            fk = slide_cfg.get("footnote_key", "footnote")
            if fk in payload:
                _fill_slide_text(
                    slide,
                    slide_cfg["footnote_shape"],
                    {0: payload[fk]},
                    font_size=9,
                    bold_token=None,
                )

        if paste_tables:
            if "picture_shape" in slide_cfg and "table_image" in slide_cfg:
                img_key = slide_cfg["table_image"]
                if img_key in images:
                    replace_picture_fit(slide, slide_cfg["picture_shape"], images[img_key])
            for pic in slide_cfg.get("pictures", []):
                if pic["table_image"] in images:
                    replace_picture_fit(slide, pic["shape"], images[pic["table_image"]])
        else:
            # 不贴新表时必须去掉模版里的旧 Q4 表图，否则看起来像「完全旧稿」
            pic_names: list[str] = []
            if "picture_shape" in slide_cfg:
                pic_names.append(slide_cfg["picture_shape"])
            for pic in slide_cfg.get("pictures", []):
                pic_names.append(pic["shape"])
            for name in pic_names:
                try:
                    _find_shape(slide, name)
                except KeyError:
                    continue
                # 清空模版旧表图，不放粘贴提示（粘贴清单见出报指引 / Excel 说明页）
                remove_picture(slide, name)

    prs.save(str(output_pptx))
    return {
        "pptx": str(output_pptx),
        "xlsx": str(xlsx_path) if xlsx_path.exists() else "",
        "work_dir": str(work_dir),
        "images": {k: str(v) for k, v in images.items()},
        "n_slides": len(slide_map.get("slides", [])),
        "table_images_from": (
            "excel" if used_excel else ("pil" if images else "skipped_manual_paste")
        ),
        "paste_tables": paste_tables,
        "narrative_source": narrative_source,
        "narrative_report": narrative_report,
    }
