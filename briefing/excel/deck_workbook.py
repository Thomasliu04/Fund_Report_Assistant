"""全册表图 Excel：各页一张 sheet，带数据条/色阶/银华高亮，便于人工调美观。"""

from __future__ import annotations

from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from briefing.excel.databar_patch import (
    patch_workbook_bidirectional_databars,
    patch_workbook_gradient_databars,
)
from briefing.excel.styled_table import ColumnSpec
from briefing.excel.table_layout import (
    excel_char_width,
    fit_excel_widths,
    font_pt_for_columns,
    format_header_wrap,
    _disp_len,
    _header_lines,
)

# 对齐 Q4 底稿：正红 / 蓝条
_BAR_RED = "FF555A"
_BAR_BLUE = "638EC6"


def _pct_column_divisor(rows: list[dict], key: str) -> float:
    """推断增速列是「百分数」(2=2%) 还是 Excel 小数 (0.02)。

    H1 底稿多为百分数（列中位数通常 ≥2）；Q4 案例多为小数（中位数 ≪1）。
    旧逻辑仅对 |v|>2 除以 100，导致 1%、2% 被写成 100%、200%。
    """
    nums = [float(r[key]) for r in rows if isinstance(r.get(key), (int, float))]
    if nums and median(abs(x) for x in nums) >= 2:
        return 100.0
    return 1.0


def _cell_value(val: Any, fmt: str, *, pct_divisor: float = 1.0):
    if val is None:
        return None
    if fmt == "int":
        try:
            return int(round(float(val)))
        except (TypeError, ValueError):
            return val
    if fmt == "pct":
        try:
            v = float(val) / pct_divisor
            return round(v, 6)
        except (TypeError, ValueError):
            return val
    return val

def _is_focus_row(company: str, focus_company: str) -> bool:
    if not focus_company or focus_company == "__none__":
        return False
    c = str(company).replace("基金", "")
    f = focus_company.replace("基金", "")
    return bool(f) and (f in c or focus_company in str(company))


def _add_sheet(
    wb: Workbook,
    title: str,
    columns: list[ColumnSpec],
    rows: list[dict],
    *,
    focus_company: str = "银华基金",
    focus_key: str = "company",
) -> tuple[list[str], list[tuple[str, str]]]:
    """写入 sheet，返回 (双向红条区域, 全部数据条区域含颜色)。"""
    ws = wb.create_sheet(title[:31])
    n_cols = len(columns)
    font_pt = font_pt_for_columns(n_cols, base=11)
    header_fill = PatternFill("solid", fgColor="F2DCDB")
    # 表头显式加粗（略大于正文字号，Excel/导出图更易辨认）
    header_font = Font(bold=True, color="000000", name="微软雅黑", size=max(font_pt, 11))
    thin = Border(
        left=Side(style="thin", color="B4C6DC"),
        right=Side(style="thin", color="B4C6DC"),
        top=Side(style="thin", color="B4C6DC"),
        bottom=Side(style="thin", color="B4C6DC"),
    )
    focus_fill = PatternFill("solid", fgColor="FCE4D6")
    body_font = Font(name="微软雅黑", size=font_pt)
    focus_font = Font(name="微软雅黑", size=font_pt, bold=True, color="9C0006")

    for c, col in enumerate(columns, 1):
        header = format_header_wrap(col.header)
        cell = ws.cell(1, c, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = thin

    pct_div = {
        col.key: _pct_column_divisor(rows, col.key) for col in columns if col.fmt == "pct"
    }

    for r_i, row in enumerate(rows, 2):
        is_focus = _is_focus_row(str(row.get(focus_key, "")), focus_company)
        for c, col in enumerate(columns, 1):
            raw = row.get(col.key)
            val = _cell_value(raw, col.fmt, pct_divisor=pct_div.get(col.key, 1.0))
            cell = ws.cell(r_i, c, val)
            cell.font = focus_font if is_focus else body_font
            cell.border = thin
            align = {"left": "left", "right": "right"}.get(col.align, "center")
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if col.fmt == "pct" and isinstance(val, (int, float)):
                cell.number_format = "0%"
            if is_focus:
                cell.fill = focus_fill

    n = len(rows) + 1
    bi_ranges: list[str] = []
    gradient_bars: list[tuple[str, str]] = []  # (sqref, rgb without alpha prefix issues)
    if n >= 2:
        for c, col in enumerate(columns, 1):
            letter = get_column_letter(c)
            rng = f"{letter}2:{letter}{n}"
            if col.bar != "none":
                bar_rgb = _BAR_BLUE if getattr(col, "bar_color", "red") == "blue" else _BAR_RED
                ws.conditional_formatting.add(
                    rng,
                    DataBarRule(
                        start_type="min",
                        end_type="max",
                        color=bar_rgb,
                        showValue=True,
                        minLength=None,
                        maxLength=None,
                    ),
                )
                gradient_bars.append((rng, bar_rgb))
                # 仅红色双向列打正负轴补丁
                if col.bar == "bidirectional" and bar_rgb == _BAR_RED:
                    bi_ranges.append(rng)
            if col.color_scale == "rank":
                ws.conditional_formatting.add(
                    rng,
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )
            elif col.color_scale == "growth":
                # 越高越红（对齐行业增速色阶）
                ws.conditional_formatting.add(
                    rng,
                    ColorScaleRule(
                        start_type="min",
                        start_color="63BE7B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="F8696B",
                    ),
                )

    # 按内容估算列宽；换行表头按最长一行保底，避免「主动权益」被裁切
    widths = [excel_char_width(col, rows, n_cols=n_cols) for col in columns]
    floors: list[float] = []
    for col in columns:
        lines = _header_lines(col.header)
        floor = 4.5
        if len(lines) >= 2:
            floor = max(_disp_len(ln) for ln in lines) + 2.4
            if len(lines[0]) >= 3:
                floor = max(floor, 9.0)
        floors.append(floor)
    widths = fit_excel_widths(widths, n_cols=n_cols, min_widths=floors)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    has_wrap = any(len(_header_lines(col.header)) >= 2 for col in columns)
    # 两行表头需要更高行高，否则第二行「排名」会被遮住
    if has_wrap:
        ws.row_dimensions[1].height = 44 if font_pt <= 10 else 48
    else:
        ws.row_dimensions[1].height = 36 if font_pt <= 10 else 32
    for r in range(2, n + 1):
        ws.row_dimensions[r].height = 17 if font_pt <= 10 else 18

    if n >= 1:
        last = get_column_letter(len(columns))
        ws.print_area = f"A1:{last}{max(n, 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.2
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    return bi_ranges, gradient_bars


def _readme_sheet(wb: Workbook, period_label: str) -> None:
    ws = wb.active
    ws.title = "使用说明"
    lines = [
        f"{period_label} 公募行业数据简报 · 表图工作簿",
        "",
        "【分工建议】",
        "· PPT：叙事文字（观点句需人工敲定）+ 贴入表图",
        "· 本 Excel：各页排名表的可编辑源文件（数据条 / 色阶 / 银华高亮）",
        "",
        "【怎么用】",
        "1. 本文件已按内容自动调整列宽；可再微调列宽/条件格式",
        "2. 生成 PPT 时会静默从本 Excel 复制表图粘贴，保证与 Excel 一致",
        "3. 也可手动：选中表格 → 复制 → 选择性粘贴到 PPT",
        "4. 数字一般不要手改；若底稿更新请重新跑生成流程",
        "",
        "【数据条】",
        "正值向右红色，负值向左绿色（与 PPT 表图一致）；轴在零点。",
        "",
        "【Sheet 对应】",
        "01_行业变化 → PPT 第1页表",
        "02_总规模 → 第2页",
        "03_非货 → 第3页",
        "04_非货增量 → 第4页",
        "05_主动权益 → 第5页",
        "06_非货ETF含联接 / 06_权益ETF含联接 → 第6页双表",
        "07_货币 → 第7页",
        "08_固收 → 第8页",
        "09_固收+ → 第9页",
        "10_FOF → 第10页",
        "",
        "【数据来源】",
        "本工作簿以底稿「终表」为准（含他表 VLOOKUP 过来的排名列），按列做数据条/色阶。",
    ]
    ws["A1"] = lines[0]
    ws["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="D32820")
    for i, line in enumerate(lines[1:], 2):
        ws.cell(i, 1, line).font = Font(name="微软雅黑", size=10)
    ws.column_dimensions["A"].width = 88


def build_deck_workbook(
    sheets: list[tuple[str, list[ColumnSpec], list[dict]]],
    output_path: str | Path,
    *,
    period_label: str = "",
    focus_company: str = "银华基金",
) -> Path:
    """
    sheets: [(sheet_title, columns, rows), ...]
    """
    wb = Workbook()
    _readme_sheet(wb, period_label or "本期")
    sheet_ranges: dict[str, list[str]] = {}
    gradient_ranges: dict[str, list[tuple[str, str]]] = {}
    for title, columns, rows in sheets:
        name = title[:31]
        bi, grads = _add_sheet(wb, name, columns, rows, focus_company=focus_company)
        if bi:
            sheet_ranges[name] = bi
        if grads:
            # 双向红条已由 bidirectional 补丁处理；此处只给蓝条/单向条补渐变
            bi_set = set(bi)
            only_grad = [(r, c) for r, c in grads if r not in bi_set]
            if only_grad:
                gradient_ranges[name] = only_grad

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    if sheet_ranges:
        patch_workbook_bidirectional_databars(out, sheet_ranges)
    if gradient_ranges:
        patch_workbook_gradient_databars(out, gradient_ranges)
    return out
