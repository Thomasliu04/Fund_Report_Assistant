"""从终表 Excel / load_final_tables 结果解析结构化事实，供 PPT 文字成文。

数字口径与贴图表一致；增速统一为小数（0.02=2%）。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from statistics import median
from typing import Any, Literal

from openpyxl import load_workbook

Severity = Literal["error", "warning", "info"]

# 终表 sheet → 逻辑名
SHEET_IDS = {
    "01_行业变化": "category",
    "02_总规模": "total",
    "03_非货": "non_money",
    "04_非货增量": "increment",
    "05_主动权益": "active_equity",
    "06_非货ETF含联接": "etf_non_money",
    "06_权益ETF含联接": "etf_equity",
    "06_ETF含联接": "etf_non_money",
    "06_ETF不含联接": "etf_equity",
    "07_货币": "money",
    "08_固收": "fixed_income",
    "09_固收+": "fixed_income_plus",
    "10_FOF": "fof",
}

CATEGORY_LABEL_TO_KEY = {
    "合计": "total",
    "非货": "non_money",
    "货币": "money",
    "固收": "fixed_income",
    "被动权益": "passive_equity",
    "主动权益": "active_equity",
    "固收+": "fixed_income_plus",
    "FOF": "fof",
}

_INC_RANK_LABELS = {
    "主动权益增量排名": "active_equity",
    "被动权益增量排名": "passive_equity",
    "固收+增量排名": "fixed_income_plus",
    "固收增量排名": "fixed_income",
    "FOF增量排名": "fof",
}


@dataclass
class FactIssue:
    severity: Severity
    code: str
    message: str
    sheet: str = ""


@dataclass
class FactRow:
    """一行公司或行业品类事实（字段按需存在）。"""

    name: str
    rank: int | None = None
    aum: float | None = None
    increment: float | None = None
    growth_pct: float | None = None  # 小数
    q_increment: float | None = None
    q_growth_pct: float | None = None
    rank_change: int | None = None  # YTD
    rank_change_q: int | None = None
    new_issue: float | None = None
    nav_change: float | None = None
    holding_sales: float | None = None
    # 增量表：增量排名 + 分项
    increment_rank: int | None = None
    category_increment_ranks: dict[str, int] = field(default_factory=dict)
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass
class SheetFacts:
    sheet_id: str
    title: str
    headers: list[str]
    rows: list[FactRow]
    issues: list[FactIssue] = field(default_factory=list)


@dataclass
class TableBook:
    sheets: dict[str, SheetFacts] = field(default_factory=dict)
    issues: list[FactIssue] = field(default_factory=list)

    def get(self, sheet_id: str) -> SheetFacts | None:
        return self.sheets.get(sheet_id)

    @property
    def errors(self) -> list[FactIssue]:
        out = [i for i in self.issues if i.severity == "error"]
        for sh in self.sheets.values():
            out.extend(i for i in sh.issues if i.severity == "error")
        return out

    @property
    def warnings(self) -> list[FactIssue]:
        out = [i for i in self.issues if i.severity == "warning"]
        for sh in self.sheets.values():
            out.extend(i for i in sh.issues if i.severity == "warning")
        return out

    def format_issues(self) -> str:
        all_i = list(self.issues)
        for sh in self.sheets.values():
            all_i.extend(sh.issues)
        if not all_i:
            return "读表检查：无问题"
        lines = [f"读表检查：error={len(self.errors)} warning={len(self.warnings)}"]
        order = {"error": 0, "warning": 1, "info": 2}
        for i in sorted(all_i, key=lambda x: (order.get(x.severity, 9), x.sheet, x.code)):
            loc = f"[{i.sheet}] " if i.sheet else ""
            lines.append(f"  {i.severity.upper():7} {i.code}: {loc}{i.message}")
        return "\n".join(lines)


def _norm_header(h: Any) -> str:
    if h is None:
        return ""
    return str(h).replace("\n", "").replace(" ", "").strip().rstrip("↓")


def _short_name(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith("基金") and len(s) > 2:
        return s[:-2]
    return s


def _to_float(v: Any) -> float | None:
    if v is None or (isinstance(v, float) and v != v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", "").replace("%", "")
    if not s or s.lower() in {"nan", "none", "-"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _to_int(v: Any) -> int | None:
    f = _to_float(v)
    if f is None:
        return None
    return int(round(f))


def _pct_divisor(values: list[float]) -> float:
    if values and median(abs(x) for x in values) >= 2:
        return 100.0
    return 1.0


def _is_pct_header(h: str) -> bool:
    return "增速" in h or "增幅" in h


def _is_ytd_window_label(h: str) -> bool:
    """前三季度 / 前两季度 / 全年 / H1 等属 YTD 窗口，不是单季。"""
    return any(
        t in h
        for t in ("前三季度", "前两季度", "前四季度", "全年", "上半年", "下半年", "YTD", "H1", "H2")
    ) or (h.endswith("年") and "季度" not in h)


def _is_single_quarter_label(h: str) -> bool:
    """单季列：季度* / Q3*；排除「前三季度」。"""
    if _is_ytd_window_label(h):
        return False
    if "季度" in h:
        return True
    if re.match(r"^Q\d", h):
        return True
    return False


def _pick_col(headers: list[str], *predicates) -> int | None:
    """返回满足任一谓词的列下标；谓词为 (norm_header) -> bool。"""
    for i, h in enumerate(headers):
        for pred in predicates:
            if pred(h):
                return i
    return None


def _col_by_equals(headers: list[str], *names: str) -> int | None:
    want = set(names)
    for i, h in enumerate(headers):
        if h in want:
            return i
    return None


def _col_contains(headers: list[str], *tokens: str, exclude: tuple[str, ...] = ()) -> int | None:
    for i, h in enumerate(headers):
        if any(ex in h for ex in exclude):
            continue
        if all(t in h for t in tokens):
            return i
    return None


def find_focus(rows: list[FactRow], focus_company: str, focus_short: str) -> FactRow | None:
    short = (focus_short or focus_company.replace("基金", "")).replace("基金", "")
    for r in rows:
        n = r.name.replace("基金", "")
        if short and short in n:
            return r
        if focus_company and focus_company in r.name:
            return r
    return None


def _map_role_indices(headers: list[str], sheet_id: str) -> dict[str, int | None]:
    """按 sheet 类型映射角色列。"""
    h = headers
    name_i = _col_by_equals(h, "类型", "公司", "基金公司", "管理人")
    if name_i is None:
        name_i = _pick_col(h, lambda x: x in {"类型", "公司", "基金公司", "管理人"})

    rank_i = _pick_col(
        h,
        lambda x: x == "排名" or (x.endswith("排名") and "变化" not in x and "增量" not in x),
    )
    aum_i = _pick_col(
        h,
        lambda x: x in {"总规模", "非货总计", "非货规模", "规模", "货币规模", "期末"},
        lambda x: x == "规模" or (x.endswith("规模") and "增量" not in x and "增幅" not in x and "上年末" not in x),
    )
    # 行业表：期末规模 = 「期末」或第一个 YYYYMMDD 列（旧表头）
    if sheet_id == "category":
        end_i = _col_by_equals(h, "期末")
        date_cols = [i for i, x in enumerate(h) if re.match(r"^\d{8}", x)]
        if end_i is not None:
            aum_i = end_i
        elif date_cols:
            aum_i = date_cols[0]

    ytd_inc = _pick_col(
        h,
        lambda x: (
            "增量" in x
            and "排名" not in x
            and not _is_single_quarter_label(x)
            and (
                _is_ytd_window_label(x)
                or "非货增量" in x
                or x in {"规模增量", "YTD非货增量", "26H1增量", "YTD增量", "YTD规模增量"}
                or x.endswith("年增量")
            )
        ),
        # 业务表裸「规模增量」；行业/总规模等非增量表可用更宽匹配
        lambda x: (
            sheet_id != "increment"
            and "增量" in x
            and "排名" not in x
            and not _is_single_quarter_label(x)
        )
        or x in {"规模增量", "YTD非货增量", "26H1增量", "YTD增量", "YTD规模增量"},
    )
    # Prefer explicit 规模增量 for business
    bi = _col_by_equals(h, "规模增量")
    if bi is not None:
        ytd_inc = bi
    yi2 = None
    for i, x in enumerate(h):
        if "非货增量" in x and "排名" not in x and not _is_single_quarter_label(x):
            yi2 = i
            break
    if sheet_id == "increment" and yi2 is not None:
        ytd_inc = yi2

    ytd_g = _pick_col(
        h,
        lambda x: _is_pct_header(x) and not _is_single_quarter_label(x),
    )
    q_inc = _pick_col(
        h,
        lambda x: _is_single_quarter_label(x) and "增量" in x and "排名" not in x,
    )
    q_g = _pick_col(
        h,
        lambda x: _is_single_quarter_label(x) and _is_pct_header(x),
    )
    rc_y = _pick_col(
        h,
        lambda x: "排名变化" in x and not _is_single_quarter_label(x),
        lambda x: x in {"排名变化", "YTD排名变化"},
    )
    # Prefer YTD 排名变化 when both exist
    ytd_rc = _col_contains(h, "YTD", "排名变化") or _col_by_equals(h, "YTD排名变化")
    if ytd_rc is not None:
        rc_y = ytd_rc
    if sheet_id in {"active_equity", "money", "fixed_income", "fixed_income_plus", "fof", "etf_non_money", "etf_equity"}:
        plain_rc = _col_by_equals(h, "排名变化")
        if plain_rc is not None:
            rc_y = plain_rc

    rc_q = _pick_col(
        h,
        lambda x: _is_single_quarter_label(x) and "排名变化" in x,
    )
    new_i = _col_by_equals(h, "新发") or _col_contains(h, "新发")
    nav_i = _pick_col(h, lambda x: "净值" in x)
    hold_i = _pick_col(h, lambda x: "持营" in x)
    inc_rank = _pick_col(
        h,
        lambda x: "非货增量" in x and "排名" in x,
        lambda x: x.endswith("增量排名") and "非货" in x,
    )

    return {
        "name": name_i,
        "rank": rank_i,
        "aum": aum_i,
        "increment": ytd_inc,
        "growth_pct": ytd_g,
        "q_increment": q_inc,
        "q_growth_pct": q_g,
        "rank_change": rc_y,
        "rank_change_q": rc_q,
        "new_issue": new_i,
        "nav_change": nav_i,
        "holding_sales": hold_i,
        "increment_rank": inc_rank,
    }


def _parse_rows(
    headers: list[str],
    raw_rows: list[list[Any]],
    sheet_id: str,
    title: str,
) -> SheetFacts:
    roles = _map_role_indices(headers, sheet_id)
    issues: list[FactIssue] = []

    if roles["name"] is None:
        issues.append(FactIssue("error", "missing_name_col", "缺少公司/类型列", title))

    # 预扫增速列，决定除数
    pct_div: dict[int, float] = {}
    for i, h in enumerate(headers):
        if _is_pct_header(h):
            vals = []
            for row in raw_rows:
                if i < len(row):
                    f = _to_float(row[i])
                    if f is not None:
                        vals.append(f)
            pct_div[i] = _pct_divisor(vals)

    def cell(row: list[Any], idx: int | None) -> Any:
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    def as_pct(row: list[Any], idx: int | None) -> float | None:
        if idx is None:
            return None
        f = _to_float(cell(row, idx))
        if f is None:
            return None
        return f / pct_div.get(idx, 1.0)

    # 增量排名分项列
    cat_rank_cols: dict[str, int] = {}
    for i, h in enumerate(headers):
        for label, key in _INC_RANK_LABELS.items():
            if h.replace("\n", "") == label or h == label:
                cat_rank_cols[key] = i
            elif "增量排名" in h and label.replace("增量排名", "") in h:
                cat_rank_cols[key] = i

    rows: list[FactRow] = []
    for row in raw_rows:
        name = _short_name(cell(row, roles["name"]))
        if not name:
            continue
        if name in {"排名", "类型"}:
            continue
        fr = FactRow(
            name=name,
            rank=_to_int(cell(row, roles["rank"])),
            aum=_to_float(cell(row, roles["aum"])),
            increment=_to_float(cell(row, roles["increment"])),
            growth_pct=as_pct(row, roles["growth_pct"]),
            q_increment=_to_float(cell(row, roles["q_increment"])),
            q_growth_pct=as_pct(row, roles["q_growth_pct"]),
            rank_change=_to_int(cell(row, roles["rank_change"])),
            rank_change_q=_to_int(cell(row, roles["rank_change_q"])),
            new_issue=_to_float(cell(row, roles["new_issue"])),
            nav_change=_to_float(cell(row, roles["nav_change"])),
            holding_sales=_to_float(cell(row, roles["holding_sales"])),
            increment_rank=_to_int(cell(row, roles["increment_rank"])),
        )
        for key, ci in cat_rank_cols.items():
            rv = _to_int(cell(row, ci))
            if rv is not None:
                fr.category_increment_ranks[key] = rv
        # 行业：把日期列等放进 extra
        if sheet_id == "category":
            for i, h in enumerate(headers):
                if re.match(r"^\d{8}", h) or "新发" in h or "净值" in h or "持营" in h or "增量" in h:
                    fr.extra[h] = _to_float(cell(row, i))
        rows.append(fr)

    if sheet_id != "category" and len(rows) < 10:
        issues.append(FactIssue("warning", "few_rows", f"仅 {len(rows)} 行数据", title))

    return SheetFacts(sheet_id=sheet_id, title=title, headers=headers, rows=rows, issues=issues)


def load_table_book_from_xlsx(xlsx_path: str | Path) -> TableBook:
    """从已导出的 xx_tables.xlsx 读表。"""
    path = Path(xlsx_path)
    book = TableBook()
    if not path.exists():
        book.issues.append(FactIssue("error", "file_missing", f"文件不存在：{path}"))
        return book

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        for title in wb.sheetnames:
            if title == "使用说明":
                continue
            sheet_id = SHEET_IDS.get(title)
            if not sheet_id:
                continue
            ws = wb[title]
            headers_raw = [ws.cell(1, c).value for c in range(1, (ws.max_column or 0) + 1)]
            # trim trailing empties
            while headers_raw and headers_raw[-1] is None:
                headers_raw.pop()
            headers = [_norm_header(h) for h in headers_raw]
            raw_rows: list[list[Any]] = []
            for r in range(2, (ws.max_row or 1) + 1):
                row = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                raw_rows.append(row)
            book.sheets[sheet_id] = _parse_rows(headers, raw_rows, sheet_id, title)
    finally:
        wb.close()

    _post_check(book)
    return book


def load_table_book_from_final_sheets(
    sheets: list[tuple[str, list, list[dict]]],
) -> TableBook:
    """从 load_final_tables 的 (title, columns, rows) 读表。"""
    book = TableBook()
    for title, columns, rows in sheets:
        sheet_id = SHEET_IDS.get(title)
        if not sheet_id:
            continue
        headers = [_norm_header(c.header) for c in columns]
        raw_rows: list[list[Any]] = []
        for row in rows:
            raw_rows.append([row.get(c.key) for c in columns])
        book.sheets[sheet_id] = _parse_rows(headers, raw_rows, sheet_id, title)
    _post_check(book)
    return book


def _post_check(book: TableBook) -> None:
    required = ("category", "total", "non_money")
    for rid in required:
        if rid not in book.sheets:
            book.issues.append(
                FactIssue("error", "missing_sheet", f"缺少终表 sheet（逻辑名 {rid}）")
            )
    cat = book.get("category")
    if cat:
        names = {r.name for r in cat.rows}
        for lab in ("合计", "非货"):
            if lab not in names:
                cat.issues.append(
                    FactIssue("error", "missing_category", f"行业表缺少「{lab}」", cat.title)
                )
    for sid in ("total", "non_money", "active_equity", "money", "fixed_income", "fof"):
        sh = book.get(sid)
        if sh and not sh.rows:
            sh.issues.append(FactIssue("error", "empty_sheet", "表无数据行", sh.title))


def require_focus(book: TableBook, sheet_id: str, focus: str, short: str) -> FactRow | None:
    sh = book.get(sheet_id)
    if not sh:
        return None
    hit = find_focus(sh.rows, focus, short)
    if not hit:
        sh.issues.append(
            FactIssue(
                "error",
                "focus_missing",
                f"未找到焦点公司「{short or focus}」",
                sh.title,
            )
        )
    return hit
