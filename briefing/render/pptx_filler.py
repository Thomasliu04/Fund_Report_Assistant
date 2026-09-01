"""将 payload 灌入 PPT 模版：整段替换文字 + 替换表图。

标题样式对齐案例简报：章节标题加粗，业务关键字标红（C00000）；
正文不再对「示例」单独加粗。
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

# 案例简报标题关键字红
_TITLE_RED = RGBColor(0xC0, 0x00, 0x00)

# 章节标题中需标红的关键字（长词优先）
_SECTION_RED_KEYWORDS = (
    "总规模（含货币）",
    "非货增量",
    "非货",
)

# 整段标红加粗的品类引导语
_LABEL_RE = re.compile(
    r"^(主动权益|货币|固收\+|固收|FOF|非货ETF（含联接）|权益ETF（含联接）)[:：]$"
)

# ETF 段首引导（冒号前标红）
_ETF_LEAD_RE = re.compile(r"^(非货ETF（含联接）|权益ETF（含联接）)[:：]")


def load_slide_map(path: str | Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _find_shape(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"未找到形状: {name!r}")


def _clear_runs(paragraph) -> None:
    """清空段落全部 run，避免残留空 run 干扰样式。"""
    p = paragraph._p
    for child in list(p):
        if child.tag.endswith("}r"):
            p.remove(child)


def _style_segments(text: str) -> list[tuple[str, bool, RGBColor | None]]:
    """
    按案例风格拆分段落为 (文本, 加粗, 颜色)。
    颜色 None = 保持模版默认（通常黑/主题色）。
    """
    if text is None:
        return [("", False, None)]
    t = text
    stripped = t.strip()

    # 封面大标题 / 数据来源等：整段加粗
    if "公募行业数据简报" in stripped and not stripped.startswith(("一、", "二、")):
        return [(t, True, None)]
    if stripped.startswith("如无特别注明"):
        return [(t, True, None)]

    # 章节标题：一、二、… 整段加粗，关键字标红
    if re.match(r"^[一二三四五六七八九十]+、", stripped):
        for kw in _SECTION_RED_KEYWORDS:
            i = t.find(kw)
            if i >= 0:
                parts: list[tuple[str, bool, RGBColor | None]] = []
                if i > 0:
                    parts.append((t[:i], True, None))
                parts.append((kw, True, _TITLE_RED))
                if i + len(kw) < len(t):
                    parts.append((t[i + len(kw) :], True, None))
                return parts
        return [(t, True, None)]

    # 品类标签整段红粗：货币： / 固收： / 主动权益：
    if _LABEL_RE.match(stripped):
        return [(t, True, _TITLE_RED)]

    # ETF 叙述：冒号及之前标红加粗，其后正文不加粗
    m = _ETF_LEAD_RE.match(stripped)
    if m:
        lead = m.group(0)
        # 保持原串前缀空白
        lead_start = t.find(lead[0]) if lead else 0
        # 更稳：在原 text 里定位 lead
        idx = t.find(lead)
        if idx < 0:
            idx = 0
            lead_len = len(lead)
        else:
            lead_len = len(lead)
        head = t[: idx + lead_len]
        tail = t[idx + lead_len :]
        parts = [(head, True, _TITLE_RED)]
        if tail:
            parts.append((tail, False, None))
        return parts

    # 小标题如「1 主动权益排名」「3.货币排名情况」：加粗、不标红
    if re.match(r"^(\d+\.\d*|\d+\.?)\s*.*排名", stripped) or re.match(
        r"^\d+\.\s*.*排名情况", stripped
    ):
        return [(t, True, None)]

    # 正文
    return [(t, False, None)]


def _set_paragraph_text(
    paragraph,
    text: str,
    *,
    font_size_pt: float | None = None,
    bold_token: str | None = None,  # 兼容旧参数；已不使用（示例不加粗）
) -> None:
    """整段替换，并按案例规则设置加粗/标红。"""
    del bold_token  # 明确不再对关键字公司加粗
    segments = _style_segments(text if text is not None else "")

    # 读取原段落基准字体（若有）
    base_size = None
    base_name = None
    if paragraph.runs:
        first = paragraph.runs[0]
        base_size = first.font.size
        base_name = first.font.name

    _clear_runs(paragraph)

    for seg, is_bold, color in segments:
        if seg == "" and len(segments) > 1:
            continue
        run = paragraph.add_run()
        run.text = seg
        if base_name:
            run.font.name = base_name
        size = Pt(font_size_pt) if font_size_pt else base_size
        if size:
            run.font.size = size
        run.font.bold = bool(is_bold)
        if color is not None:
            run.font.color.rgb = color


def replace_narrative_paragraphs(
    shape,
    paragraph_map: dict[int, str],
    font_size_pt: float = 12,
    *,
    bold_token: str | None = None,
) -> None:
    """按段落索引写入文本。paragraph_map: {para_index: text}"""
    paragraphs = shape.text_frame.paragraphs
    for idx, text in paragraph_map.items():
        if idx >= len(paragraphs):
            continue
        if text is None:
            continue
        _set_paragraph_text(
            paragraphs[idx], text, font_size_pt=font_size_pt, bold_token=bold_token
        )


def remove_picture(slide, shape_name: str) -> bool:
    """删除幻灯片上指定名称的图片（用于去掉模版残留的旧表图）。"""
    try:
        shape = _find_shape(slide, shape_name)
    except KeyError:
        return False
    sp = shape._element
    sp.getparent().remove(sp)
    return True


def add_table_placeholder(
    slide,
    *,
    left,
    top,
    width,
    height,
    text: str = "【请从本期 tables.xlsx 对应 sheet 复制表图粘贴至此】",
) -> None:
    """在原图表位置放提示框，避免空白页不知道该贴哪里。"""
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = "表图占位提示"
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    if p.runs:
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.bold = True


def replace_picture_fit(
    slide,
    shape_name: str,
    image_path: str | Path,
    *,
    prefer: str = "width",
    valign: str = "top",
) -> None:
    """
    替换图片并保持宽高比放入原图框。

    prefer:
      - width: 优先铺满框宽（表格更易读），过高再按高度缩小
      - height: 优先铺满框高
      - contain: 完整落入框内（旧行为，居中）
    valign: top | center — 垂直对齐
    """
    shape = _find_shape(slide, shape_name)
    box_left, box_top = shape.left, shape.top
    box_w, box_h = shape.width, shape.height
    sp = shape._element
    sp.getparent().remove(sp)

    from PIL import Image

    with Image.open(image_path) as im:
        img_w, img_h = im.size
    if img_w <= 0 or img_h <= 0:
        return
    img_aspect = img_w / img_h

    if prefer == "width":
        new_w = box_w
        new_h = int(box_w / img_aspect)
        if new_h > box_h:
            new_h = box_h
            new_w = int(box_h * img_aspect)
    elif prefer == "height":
        new_h = box_h
        new_w = int(box_h * img_aspect)
        if new_w > box_w:
            new_w = box_w
            new_h = int(box_w / img_aspect)
    else:
        box_aspect = box_w / box_h
        if img_aspect > box_aspect:
            new_w = box_w
            new_h = int(box_w / img_aspect)
        else:
            new_h = box_h
            new_w = int(box_h * img_aspect)

    left = box_left + (box_w - new_w) // 2
    if valign == "center":
        top = box_top + (box_h - new_h) // 2
    else:
        top = box_top
    slide.shapes.add_picture(str(image_path), left, top, width=new_w, height=new_h)


def replace_picture_inplace(slide, shape_name: str, image_path: str | Path, *, fit: bool = True) -> None:
    if fit:
        replace_picture_fit(slide, shape_name, image_path)
        return
    shape = _find_shape(slide, shape_name)
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    sp = shape._element
    sp.getparent().remove(sp)
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def fill_pptx(
    template_path: str | Path,
    slide_map: dict,
    payload: dict[str, Any],
    output_path: str | Path,
    image_paths: dict[str, Path],
    *,
    bold_token: str | None = None,
) -> Path:
    """
    payload 结构示例::

        {
          "overview": {
            "title": "...",
            "data_source_note": "...",
            ...
            "footnote": "..."
          },
          "total_ranking": { ... }
        }

    image_paths: {"category_summary": Path(...), "total_ranking_top30": Path(...)}
    """
    template_path = Path(template_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template_path, output_path)

    prs = Presentation(str(output_path))

    for slide_cfg in slide_map.get("slides", []):
        slide = prs.slides[slide_cfg["index"]]
        section_id = slide_cfg["id"]
        section_payload = payload.get(section_id, {})

        if "narrative_shape" in slide_cfg and "paragraphs" in slide_cfg:
            shape = _find_shape(slide, slide_cfg["narrative_shape"])
            para_map = {}
            for para_idx, key in slide_cfg["paragraphs"].items():
                idx = int(para_idx)
                if key in section_payload:
                    para_map[idx] = section_payload[key]
            replace_narrative_paragraphs(
                shape, para_map, font_size_pt=12, bold_token=bold_token
            )

        if "footnote_shape" in slide_cfg:
            fk = slide_cfg.get("footnote_key", "footnote")
            if fk in section_payload:
                fn_shape = _find_shape(slide, slide_cfg["footnote_shape"])
                replace_narrative_paragraphs(
                    fn_shape, {0: section_payload[fk]}, font_size_pt=10, bold_token=None
                )

        table_key = slide_cfg.get("table_image")
        if table_key and table_key in image_paths and "picture_shape" in slide_cfg:
            replace_picture_inplace(slide, slide_cfg["picture_shape"], image_paths[table_key])

    prs.save(str(output_path))
    return output_path
